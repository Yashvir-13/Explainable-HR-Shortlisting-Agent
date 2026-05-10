from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SLIDES = [
    ("Problem", "Manual screening creates fatigue, inconsistency, and bias risk.", ["Hundreds of applicants per role", "Scorecards vary by reviewer", "Audit trail is often thin"]),
    ("Solution", "A transparent agent ranks candidates while keeping HR in control.", ["JD and profile parsing", "Five-dimension rubric", "Override reason captured"]),
    ("Agent Flow", "A bounded LangGraph pipeline makes the workflow explainable.", ["Load and sanitize", "Parse JD and profiles", "Score, rank, report"]),
    ("Architecture", "Groq reasoning is optional; deterministic fallback keeps demos reliable.", ["Groq JSON mode for structured extraction", "Pydantic validation", "HTML and JSON reports"]),
    ("Rubric", "Every candidate receives the same weighted evaluation.", ["Skills Match: 30%", "Experience Relevance: 25%", "Education & Certs: 15%", "Project / Portfolio: 20%", "Communication Quality: 10%"]),
    ("Security", "The prototype treats resumes as untrusted and sensitive inputs.", ["Prompt-injection cleanup", "PII masking before cloud calls", ".env-based secrets", "Human review gate"]),
    ("Human Review", "The agent recommends; HR decides.", ["Dimension-level justifications", "Required reason for overrides", "Append-only audit log"]),
    ("Demo Data", "The sample batch validates strong, partial, adjacent, and weak matches.", ["Five sample LinkedIn JSON profiles", "Senior AI HR role", "Generated shortlist report"]),
    ("Results", "The best-fit candidate rises because the evidence aligns with the JD.", ["Isha: hire", "Rahul and Neha: hold", "Omar and Vikram: no-hire"]),
    ("Learnings", "Structured outputs and scoped automation make the system easier to trust.", ["Start with schemas", "Cache and trace during development", "Keep humans in the loop"]),
]


def add_textbox(slide, left, top, width, height, text, size=22, bold=False, color=RGBColor(26, 32, 44)):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def build_deck(output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    ink = RGBColor(23, 32, 42)
    accent = RGBColor(22, 119, 255)
    muted = RGBColor(91, 102, 117)
    paper = RGBColor(248, 250, 252)

    for index, (kicker, claim, bullets) in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = paper

        add_textbox(slide, Inches(0.65), Inches(0.45), Inches(2.5), Inches(0.3), f"{index:02d} / 10", 10, True, muted)
        add_textbox(slide, Inches(0.65), Inches(0.82), Inches(3.8), Inches(0.35), kicker.upper(), 12, True, accent)
        add_textbox(slide, Inches(0.65), Inches(1.35), Inches(8.4), Inches(1.2), claim, 34, True, ink)

        top = Inches(3.0)
        for bullet in bullets:
            shape = slide.shapes.add_shape(1, Inches(0.72), top + Inches(0.12), Inches(0.09), Inches(0.09))
            shape.fill.solid()
            shape.fill.fore_color.rgb = accent
            shape.line.color.rgb = accent
            add_textbox(slide, Inches(0.95), top, Inches(7.2), Inches(0.35), bullet, 18, False, ink)
            top += Inches(0.56)

        panel = slide.shapes.add_shape(1, Inches(9.4), Inches(1.1), Inches(3.05), Inches(5.35))
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(235, 241, 247)
        panel.line.color.rgb = RGBColor(210, 221, 232)
        add_textbox(slide, Inches(9.75), Inches(1.55), Inches(2.25), Inches(0.45), "Prototype", 14, True, muted)
        add_textbox(slide, Inches(9.75), Inches(2.18), Inches(2.3), Inches(1.25), "Explainable scoring with audit-ready evidence.", 23, True, ink)
        add_textbox(slide, Inches(9.75), Inches(4.55), Inches(2.25), Inches(0.9), "LangGraph + Groq + Streamlit", 16, False, accent)

        footer = add_textbox(slide, Inches(0.65), Inches(6.95), Inches(8.0), Inches(0.25), "Explainable HR Shortlisting Agent", 9, False, muted)
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


if __name__ == "__main__":
    build_deck(Path("outputs/explainable_hr_shortlisting_agent_deck.pptx"))

